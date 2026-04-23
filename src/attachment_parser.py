import io
import logging

logging.getLogger("pdfminer").setLevel(logging.ERROR)


def parse_pdf(data: bytes) -> str:
    import pdfplumber
    text = ""
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return text


def parse_excel(data: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    text = ""
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        text += f"[シート: {sheet_name}]\n"
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            row_text = " | ".join(cells).strip(" |")
            if row_text:
                text += row_text + "\n"
    return text


def parse_attachment(filename: str, data: bytes) -> str:
    name = filename.lower()
    try:
        if name.endswith(".pdf"):
            return parse_pdf(data)
        elif name.endswith((".xlsx", ".xls")):
            return parse_excel(data)
    except Exception as e:
        return f"[添付ファイル解析エラー: {e}]"
    return ""
